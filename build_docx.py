# -*- coding: utf-8 -*-
"""Markdown -> Word 转换器（常驻工具，勿删）

关键设计：
  1. 编号列表输出**字面编号**（普通段落 + 悬挂缩进），不用 Word 的
     `List Number` 样式。原因：Word 会把全文所有 List Number 段落串成
     一条编号序列，导致 Section 6 的列表从 4 开始而非 1。
  2. 表格解析跳过 markdown 分隔行 `| --- |`，并对重复表头去重。
  3. 中韩文字符显式设置 w:eastAsia 字体，否则 Word 中显示为方框。

用法:  <python> build_docx.py
"""
import os
import re
import shutil
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.stdout.reconfigure(encoding="utf-8")
SP = r"D:\yanjiubaogaoxiangmu\sci_paper"
FIGDIR = os.path.join(SP, "figures")
LATIN, EAST, MONO = "Times New Roman", "Malgun Gothic", "Consolas"
CJK = re.compile(r"[\u3000-\u9fff\uac00-\ud7af\uff00-\uffef]")
TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")


def sr(run, bold=None, italic=None, mono=False, size=None):
    run.font.name = MONO if mono else LATIN
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rf)
    rf.set(qn("w:ascii"), MONO if mono else LATIN)
    rf.set(qn("w:hAnsi"), MONO if mono else LATIN)
    if CJK.search(run.text or ""):
        rf.set(qn("w:eastAsia"), EAST)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if size:
        run.font.size = Pt(size)


def rich(par, text, size=10.5):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            sr(par.add_run(part[2:-2]), bold=True, size=size)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            sr(par.add_run(part[1:-1]), mono=True, size=size)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            sr(par.add_run(part[1:-1]), italic=True, size=size)
        else:
            sr(par.add_run(part), size=size)


def setup(doc):
    st = doc.styles["Normal"]
    st.font.name, st.font.size = LATIN, Pt(10.5)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rf)
    rf.set(qn("w:ascii"), LATIN)
    rf.set(qn("w:hAnsi"), LATIN)
    rf.set(qn("w:eastAsia"), EAST)
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.5)


def isrow(l):
    l = l.strip()
    return l.startswith("|") and l.endswith("|")


def is_sep_row(l):
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", l.strip()))


def cells(l):
    return [c.strip() for c in l.strip().strip("|").split("|")]


def add_table(doc, rows):
    head, body = rows[0], rows[1:]
    if body and body[0] == head:
        body = body[1:]
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        c.text = ""
        rich(c.paragraphs[0], h, 9)
        for r in c.paragraphs[0].runs:
            r.bold = True
    for row in body:
        cs = t.add_row().cells
        for i, v in enumerate(row[:len(head)]):
            cs[i].text = ""
            rich(cs[i].paragraphs[0], v, 9)
    doc.add_paragraph()


FIG = {
    "### 3.3": [("fig1_tokenization.png", "Figure 1. Whitespace-tokenization statistics for the three MASSIVE locales. The Chinese bar is the defect: 97.6% of utterances become a single token."),
                ("fig2_vocabulary.png", "Figure 2. Whitespace vocabulary size and hapax ratio. The Chinese vocabulary is 11,142 entries of which 96.3% occur exactly once - the tokenizer treats each sentence as one opaque symbol.")],
    "### 4.1": [("fig3_accuracy.png", "Figure 3. In-language accuracy on MASSIVE under the default (whitespace) condition.")],
    "### 4.2": [("fig4_inflation.png", "Figure 4. The four train-source x test-source cells under character-level segmentation, on the fixed evaluation subset. Dark bars are the localized-trained model at localized and translated test; light bars are the translated-trained model at translated and localized test. The two mismatches are not symmetric, and the sign of the effect reverses between the two starting points.")],
    "### 4.3": [("fig5_dissociation.png", "Figure 5. The double dissociation. Masking the padding position repairs English and Korean but not Chinese, isolating tokenization from the padding defect."),
                ("fig6_tokenization_effect.png", "Figure 6. Effect of tokenization on 60-class intent classification. Panel (a) TextCNN, panel (b) BiLSTM.")],
}


def conv(md):
    lines = open(os.path.join(SP, md), encoding="utf-8").read().split("\n")
    doc = Document()
    setup(doc)
    i, first = 0, False
    while i < len(lines):
        s = lines[i].strip()
        if not s or re.fullmatch(r"-{3,}", s):
            i += 1
            continue
        if isrow(s):
            raw = []
            while i < len(lines) and isrow(lines[i].strip()):
                raw.append(lines[i])
                i += 1
            rows = [cells(l) for l in raw if not is_sep_row(l)]
            if len(rows) >= 2:
                add_table(doc, rows)
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            lvl, txt = len(m.group(1)), m.group(2).strip()
            if lvl == 1 and not first:
                h = doc.add_heading("", level=0)
                rich(h, txt, 16)
                first = True
            else:
                h = doc.add_heading("", level=min(lvl - 1, 3))
                rich(h, txt, 13 if lvl <= 2 else 11.5)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0, 0, 0)
            for k, figs in FIG.items():
                if s.startswith(k):
                    for fn, cap in figs:
                        p = os.path.join(FIGDIR, fn)
                        if os.path.exists(p):
                            doc.add_picture(p, width=Cm(15.5))
                            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                            cp = doc.add_paragraph()
                            rich(cp, cap, 9)
                            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            rich(p, " ".join(x for x in buf if x), 10)
            for r in p.runs:
                r.italic = True
            continue
        if re.match(r"^[-*]\s+", s):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.75)
            rich(p, "\u2022  " + re.sub(r"^[-*]\s+", "", s))
            i += 1
            continue
        mo = re.match(r"^(\d+)[.)]\s+(.*)$", s)
        if mo:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.9)
            p.paragraph_format.first_line_indent = Cm(-0.9)
            rich(p, mo.group(1) + ".  " + mo.group(2))
            i += 1
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        rich(p, s)
        i += 1
    return doc


def save_smart(doc, dx):
    target = os.path.join(SP, dx)
    tmp = target + ".tmp"
    doc.save(tmp)
    try:
        os.replace(tmp, target)
        return target, "已覆盖"
    except PermissionError:
        alt = target.replace(".docx", "_revised.docx")
        try:
            os.replace(tmp, alt)
        except PermissionError:
            alt = target.replace(".docx", "_revised2.docx")
            shutil.move(tmp, alt)
        return alt, "原文件被占用 -> 另存修订版"


if __name__ == "__main__":
    for md, dx in [("manuscript_draft.md", "Manuscript_What_a_Benchmark_Number_Actually_Measures.docx"),
                   ("SUBMISSION_PACKAGE.md", "Submission_Package.docx")]:
        d = conv(md)
        path, note = save_smart(d, dx)
        print("%-62s %7.0f KB   %s" % (os.path.basename(path), os.path.getsize(path) / 1024, note))
