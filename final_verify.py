# -*- coding: utf-8 -*-
"""最终总核对：全部问题是否已解决、是否已进入最新版论文。"""
import os
import re
import sys
import zipfile

sys.stdout.reconfigure(encoding="utf-8")
SP = r"D:\yanjiubaogaoxiangmu\sci_paper"
md = open(SP + r"\manuscript_draft.md", encoding="utf-8").read()
pkg = open(SP + r"\SUBMISSION_PACKAGE.md", encoding="utf-8").read()
docx = SP + r"\Manuscript_What_a_Benchmark_Number_Actually_Measures.docx"

# ---- Word 实际内容（用于确认"已进入最新版论文"）----
from docx import Document
d = Document(docx)
dtxt = "\n".join(p.text for p in d.paragraphs)
for tb in d.tables:
    for row in tb.rows:
        dtxt += "\n" + " | ".join(c.text for c in row.cells)
z = zipfile.ZipFile(docx)
nimg = len([n for n in z.namelist() if n.startswith("word/media/")])

def both(s):
    """Markdown 与 Word 都含该串"""
    return (s in md, s in dtxt)

R = []
def chk(group, name, ok, note=""):
    R.append((group, name, ok, note))

# ============ 一、审稿意见 54 项（抽查关键项）============
G = "审稿意见"
for nm, s in [
    ("C1 random-100 重做", "random-100 only"),
    ("C5 删掉「约两倍」", None),
    ("C7 kappa 150 条", "150 non-duplicated items"),
    ("D1 删掉 all three checks", None),
    ("D2 删掉对称论证", None),
    ("E2 预训练区间", "+0.56, +2.93"),
    ("A1 80.60→80.55", "80.55"),
    ("A7 0.49→0.58", "0.58 points in English"),
    ("A14 native subword", "native subword"),
    ("A16 小节重编号", "#### 4.2.6"),
    ("A17 官方划分", "no split seed to vary"),
    ("A18 Appendix B 已除", None),
    ("B1 摘要小标题", "weakens with model strength"),
    ("B2 三分之一", "about a third on Chinese"),
    ("B6 chrF 措辞", "not explained by sentence length on either side"),
    ("F5 Xiaoya Li 2019", "Is Word Segmentation Necessary"),
]:
    if s is None:
        continue
    a, b = both(s)
    chk(G, nm, a and b, "" if (a and b) else ("md=%s word=%s" % (a, b)))

for nm, s in [("C5 无「约两倍」", "twice as likely"), ("D1 无 all three checks", "All three checks pass"),
              ("D2 无 symmetrically", "symmetrically"), ("A18 无 Appendix B", "Appendix B"),
              ("无 unverified", "unverified")]:
    a, b = both(s)
    chk(G, nm, (not a) and (not b), "" if (not a and not b) else ("md=%s word=%s" % (a, b)))

# ============ 二、三大新增实验 ============
G = "新增实验"
for nm, s in [
    ("v4a/v4c 预训练 2×2", "Table 5g"),
    ("预训练行数据", "−17.58"),
    ("模型强度单调", "shrinks monotonically"),
    ("v4b 跨系统", "Table 5h"),
    ("v4b 数值", "+10.11"),
    ("v4b 缩减幅度", "9 to 16%"),
    ("P1 人工评判", "Table 11"),
    ("P1 主结论", "1.35–1.52"),
    ("P1 随机样本", "71–106%"),
    ("审计四项", "4.34 ± 0.20"),
    ("审计路径 MD5", "135c741954f86f23d37784dba247c78a"),
    ("审计逐意图", "37 of the 59 intents exceed 50%"),
    ("审计剥离", "66.48 / 63.69 / 66.71"),
    ("chrF 译文长度", "+1.88"),
]:
    a, b = both(s)
    chk(G, nm, a and b, "" if (a and b) else ("md=%s word=%s" % (a, b)))

# ============ 三、结构与格式 ============
G = "结构"
hs = [l for l in md.split("\n") if l.startswith("#")]
chk(G, "标题层级完整", len(hs) >= 38, "%d 个" % len(hs))
chk(G, "Figure 1–6 全在正文引用", all(("Figure %d" % i) in md for i in range(1, 7)))
chk(G, "Word 含 6 张图", nimg == 6, "%d 张" % nimg)
chk(G, "无残留转义竖线", "\\|" not in md)
chk(G, "无 ASCII 箭头", "->" not in md)
chk(G, "拼写统一 tokenization", md.count("tokenisation") == 0)
rt = 1000 * len(re.findall("rather than", md)) / len(md.split())
chk(G, "rather than 密度 <2.0", rt < 2.0, "%.2f/千词" % rt)
chk(G, "无重复顶层标题", len([h for h in hs if h.startswith("# ")]) == 1)
chk(G, "§2.1 位于 §2 之后", md.find("### 2.1") > md.find("## 2. Related Work"))

# ============ 四、投稿包 ============
G = "投稿包"
for nm, s in [("IRB 含 4 位标注者", "Four human annotators"), ("知情同意", "Informed consent was obtained"),
              ("CRediT 标注分工", "Annotation (Section 4.6)"), ("run 数 353", "353 de-duplicated"),
              ("v4b 清单", "runs_v4b"), ("P1 清单", "p1_judging"),
              ("审计清单", "audit_zh_bert_v2"), ("无 unverified", None)]:
    if s is None:
        chk(G, "投稿包无 unverified", "unverified" not in pkg)
    else:
        chk(G, nm, s in pkg)

# ============ 输出 ============
ok = [r for r in R if r[2]]
bad = [r for r in R if not r[2]]
cur = None
for g, n, o, note in R:
    if g != cur:
        print()
        print("【%s】" % g)
        cur = g
    print("  %s %-34s %s" % ("OK " if o else "!! ", n, note))
print()
print("=" * 90)
print("  通过 %d / %d" % (len(ok), len(R)))
if bad:
    print("  **未通过 %d 项：**" % len(bad))
    for g, n, o, note in bad:
        print("     %s - %s %s" % (g, n, note))
print("=" * 90)
print()
print("  论文规模: %d 词 / %d 表 / 6 图 / %d 条参考文献" % (
    len(md.split()), len(re.findall(r"\*\*Table \d", md)), md.count("\n• ")))
print("  Word 段落 %d / 表格 %d / 图片 %d" % (len(d.paragraphs), len(d.tables), nimg))
